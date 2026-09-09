"""
Generador de reportes en Excel (.xlsx) de negocios prospectados.

Incluye, por cada negocio: datos de contacto, la descripción/giro
extraída del scraping, el campo de notas (donde el vendedor anota
manualmente lo que encontró al visitar -- menú, precios, promociones,
etc., ya que Google Maps no expone esa información de forma estructurada
en los resultados de búsqueda), y el pipeline de ventas completo: etapa,
café ofrecido, gramos de muestra dejados, fecha de última visita, si se
volvió cliente, cuánto consume, distancia y costo estimado de visita
(ida y vuelta) desde el punto de referencia fijo.
"""
from __future__ import annotations

import io
from typing import Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.core.config import settings
from app.models.business import Business
from app.services.coffee_catalog import calculate_sample_cost
from app.services.reference_point_service import ReferencePoint
from app.services.route_planner import haversine_km

_STAGE_LABELS = {
    "nuevo": "Nuevo",
    "contactado": "Contactado",
    "visitado": "Visitado",
    "cliente": "Cliente",
    "descartado": "Descartado",
}

_FREQUENCY_LABELS = {
    "semanal": "Semanal",
    "quincenal": "Quincenal",
    "mensual": "Mensual",
}

_HEADERS = [
    "Nombre",
    "Categoría",
    "Zona",
    "Distancia (km)",
    "Costo estimado de visita (ida y vuelta)",
    "Dirección",
    "Teléfono",
    "Sitio web",
    "Rating",
    "Etapa",
    "Descripción / Productos y servicios",
    "Café ofrecido / mostrado",
    "Muestra dejada (g)",
    "Costo real de la muestra",
    "Notas (menú, precios, etc.)",
    "Visitado",
    "Última visita",
    "¿Es cliente?",
    "Consumo",
    "Frecuencia de compra",
    "Latitud",
    "Longitud",
]

_COLUMN_WIDTHS = [28, 18, 14, 12, 20, 36, 16, 30, 8, 14, 40, 30, 14, 16, 40, 10, 16, 10, 12, 16, 12, 12]

# Índices de columna (1-based) que reciben el estilo de texto largo con
# ajuste de línea, para mantenerlos legibles sin importar el contenido.
_WRAP_TEXT_COLUMNS = (11, 12, 15)


def build_businesses_report(businesses: Sequence[Business], reference: ReferencePoint) -> bytes:
    """Construye un archivo .xlsx en memoria con una fila de contexto (el
    punto de referencia exacto desde el que se miden distancias/costos),
    una fila de encabezados, una fila por negocio, y una fila de totales
    al final (gramos de muestra, costo real de muestras, y costo de
    visita acumulados). Devuelve los bytes del archivo, listos para
    enviar como respuesta descargable.
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Prospectos"

    num_columns = len(_HEADERS)

    # --- Fila 1: contexto del punto de referencia exacto ---
    reference_note = (
        f"Distancias y costos calculados desde: {reference.label} "
        f"(lat {reference.lat}, lng {reference.lng}) "
        f"· Costo asumido: {settings.COST_PER_KM} por km "
        f"· Costo de muestras según lista de precios Nudo Verde"
    )
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=num_columns)
    note_cell = sheet.cell(row=1, column=1, value=reference_note)
    note_cell.font = Font(italic=True, color="374151")
    note_cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 24

    # --- Fila 2: encabezados ---
    header_row = 2
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1D4ED8", end_color="1D4ED8", fill_type="solid")

    for col_index, header in enumerate(_HEADERS, start=1):
        cell = sheet.cell(row=header_row, column=col_index, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    # --- Filas de datos ---
    total_sample_grams = 0.0
    total_sample_cost = 0.0
    total_visit_cost = 0.0
    data_start_row = header_row + 1

    for offset, business in enumerate(businesses):
        row_index = data_start_row + offset
        distance_km = haversine_km(reference.lat, reference.lng, business.lat, business.lng)
        visit_cost = round(distance_km * 2 * settings.COST_PER_KM, 2)
        distance_km = round(distance_km, 2)
        sample_cost = calculate_sample_cost(business.coffee_offered, business.sample_grams)

        total_visit_cost += visit_cost
        if business.sample_grams:
            total_sample_grams += business.sample_grams
        if sample_cost:
            total_sample_cost += sample_cost

        last_visited = (
            business.last_visited_at.strftime("%Y-%m-%d %H:%M")
            if business.last_visited_at
            else ""
        )
        frequency_label = (
            _FREQUENCY_LABELS.get(business.purchase_frequency.value, "")
            if business.purchase_frequency
            else ""
        )

        sheet.cell(row=row_index, column=1, value=business.name)
        sheet.cell(row=row_index, column=2, value=business.category.value)
        sheet.cell(row=row_index, column=3, value=business.zone.value)
        sheet.cell(row=row_index, column=4, value=distance_km)
        sheet.cell(row=row_index, column=5, value=visit_cost)
        sheet.cell(row=row_index, column=6, value=business.address or "")
        sheet.cell(row=row_index, column=7, value=business.phone or "")
        sheet.cell(row=row_index, column=8, value=business.website or "")
        sheet.cell(row=row_index, column=9, value=business.rating)
        sheet.cell(
            row=row_index,
            column=10,
            value=_STAGE_LABELS.get(business.stage.value, business.stage.value),
        )
        sheet.cell(row=row_index, column=11, value=business.products_services or "")
        sheet.cell(row=row_index, column=12, value=business.coffee_offered or "")
        sheet.cell(row=row_index, column=13, value=business.sample_grams)
        sheet.cell(row=row_index, column=14, value=sample_cost)
        sheet.cell(row=row_index, column=15, value=business.notes or "")
        sheet.cell(row=row_index, column=16, value="Sí" if business.visited else "No")
        sheet.cell(row=row_index, column=17, value=last_visited)
        sheet.cell(row=row_index, column=18, value="Sí" if business.is_client else "No")
        sheet.cell(row=row_index, column=19, value=business.purchase_amount)
        sheet.cell(row=row_index, column=20, value=frequency_label)
        sheet.cell(row=row_index, column=21, value=business.lat)
        sheet.cell(row=row_index, column=22, value=business.lng)

        for col_index in _WRAP_TEXT_COLUMNS:
            sheet.cell(row=row_index, column=col_index).alignment = Alignment(
                wrap_text=True, vertical="top"
            )

    # --- Fila de totales (justo después del último negocio) ---
    totals_row = data_start_row + len(businesses)
    total_label_cell = sheet.cell(row=totals_row, column=3, value="TOTALES →")
    total_label_cell.font = Font(bold=True)
    total_label_cell.alignment = Alignment(horizontal="right")

    total_cost_cell = sheet.cell(row=totals_row, column=5, value=round(total_visit_cost, 2))
    total_cost_cell.font = Font(bold=True)

    total_grams_cell = sheet.cell(row=totals_row, column=13, value=round(total_sample_grams, 1))
    total_grams_cell.font = Font(bold=True)

    total_sample_cost_cell = sheet.cell(
        row=totals_row, column=14, value=round(total_sample_cost, 2)
    )
    total_sample_cost_cell.font = Font(bold=True)

    for col_index, width in enumerate(_COLUMN_WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(col_index)].width = width

    sheet.freeze_panes = f"A{data_start_row}"

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
